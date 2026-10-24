from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_manual():
    doc = Document()
    
    # Título Principal
    title = doc.add_heading('MANUAL DE LANZAMIENTO: TIENDA ONLINE "A TU MARCA"', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Subtítulo
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Propietario: Coach Ruben Ramírez | Estatus: Listo para Ejecución')
    run.italic = True

    # Sección 1: Activos Técnicos
    doc.add_heading('1. ACTIVOS TÉCNICOS (WEB)', level=1)
    doc.add_paragraph('Archivo: countdown_3d.html', style='List Bullet')
    doc.add_paragraph('Concepto: Landing page cinematográfica con un cubo 3D interactivo.', style='List Bullet')
    doc.add_paragraph('Función: Generar hype y urgencia mediante un contador regresivo dinámico.', style='List Bullet')
    doc.add_paragraph('Acción: Subir el archivo al servidor o hosting elegido.', style='List Bullet')

    # Sección 2: Optimización de Instagram
    doc.add_heading('2. OPTIMIZACIÓN DE INSTAGRAM (CONVERSIÓN)', level=1)
    
    doc.add_heading('Configuración de Perfil', level=2)
    doc.add_paragraph('Name Field (SEO): Ruben Ramírez | Running 35+', style='List Bullet')
    doc.add_paragraph('Bio:', style='List Bullet')
    bio_p = doc.add_paragraph('Ayudo a adultos 35+ a pasar del sofá a sus primeros 5K sin lesiones. 🏃‍♂️\nIngeniería aplicada al running. ⚙️\n👇 Apertura de Tienda en 3 días\n[Link de la web]')
    bio_p.paragraph_format.left_indent = Pt(20)

    doc.add_heading('Historias Destacadas (Highlights)', level=2)
    doc.add_paragraph('Método: Explicación de la periodización y la fuerza.', style='List Bullet')
    doc.add_paragraph('Resultados: Prueba social y testimonios de atletas.', style='List Bullet')
    doc.add_paragraph('Sobre Mí: Trayectoria TI + Coaching.', style='List Bullet')
    doc.add_paragraph('Tienda: Cuenta regresiva y adelantos.', style='List Bullet')

    doc.add_heading('Estrategia de Anclados (Pinned Posts)', level=2)
    doc.add_paragraph('1. Post de Presentación (Manifiesto del Ingeniero-Coach)', style='List Bullet')
    doc.add_paragraph('2. Post de Método (Carrusel de errores a los 35+)', style='List Bullet')
    doc.add_paragraph('3. Post de Prueba Social (Testimonio de éxito)', style='List Bullet')

    # Sección 3: Plan de Contenidos
    doc.add_heading('3. PLAN DE CONTENIDOS: CUENTA REGRESIVA', level=1)
    
    doc.add_heading('Día 3: El Dolor y la Promesa (Carrusel)', level=2)
    doc.add_paragraph('Concepto: "Por qué la mayoría abandona el running a los 35".', style='List Bullet')
    doc.add_paragraph('Acción: 8 slides analizando errores y presentando el método.', style='List Bullet')
    doc.add_paragraph('CTA: Activar notificaciones.', style='List Bullet')

    doc.add_heading('Día 2: La Autoridad (Reel)', level=2)
    doc.add_paragraph('Concepto: "La Arquitectura del Running".', style='List Bullet')
    doc.add_paragraph('Acción: Contraste entre hardware (cuerpo) vs software (entrenamiento).', style='List Bullet')
    doc.add_paragraph('CTA: Abrimos en 48h. Activar notificaciones.', style='List Bullet')

    doc.add_heading('Día 1: La Anticipación Final (Stories + Post)', level=2)
    doc.add_paragraph('Acción: Publicar imagen minimalista de "MAÑANA".', style='List Bullet')
    doc.add_paragraph('Interacción: Palabra clave "LISTO" por DM para enlace.', style='List Bullet')

    # Sección 4: Copys Finales
    doc.add_heading('4. COPIES FINALES', level=1)
    
    doc.add_heading('Post de Presentación', level=2)
    doc.add_paragraph('Carrusel de 5 slides: Desde el "Por qué" hasta la misión de A Tu Marca.', style='List Bullet')
    doc.add_paragraph('Caption: Enfocado en la dualidad Ingeniería -> Running.', style='List Bullet')

    doc.add_heading('Post Día 3', level=2)
    doc.add_paragraph('Caption: "Correr después de los 35 no es un problema de voluntad, es un problema de método..."', style='List Bullet')

    doc.add_heading('Reel Día 2', level=2)
    doc.add_paragraph('Guion: Hoja de rodaje de 30s con cambios de frame.', style='List Bullet')

    # Sección 5: Gobierno
    doc.add_heading('5. GOBIERNO DEL PROYECTO (CLAUDE.md)', level=1)
    doc.add_paragraph('Técnica: Arquitectura Hexagonal y Pipeline SDD.', style='List Bullet')
    doc.add_paragraph('Calidad: Seguridad, Idempotencia y Testing obligatorio.', style='List Bullet')
    doc.add_paragraph('Tono: Empatía con franqueza, bilingüe, sin frases de gurú.', style='List Bullet')

    doc.save('Lanzamiento_IG.docx')

create_manual()
